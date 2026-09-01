# -*- coding: utf-8 -*-
"""
遥感监测图斑逐县统计脚本
生成包含原始下发、核实后、非农违法、非粮违法的逐县统计表
"""

import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
import sys
import os

def generate_county_stats(input_file, output_file):
    """
    生成逐县统计表
    
    参数:
        input_file: 输入的遥感检测检索明细表Excel文件路径
        output_file: 输出的统计表Excel文件路径
    """
    
    # 读取数据
    df = pd.read_excel(input_file)
    
    # 数据筛选
    original = df[df['图斑状态'].isna() | (df['图斑状态'] == '被拆分')]
    verified = df[df['图斑状态'].isna() | (df['图斑状态'] == '拆分')]
    illegal = verified[verified['核实认定意见'].str.contains('非农违法', na=False)]
    grain_illegal = verified[verified['核实认定意见'].str.contains('非粮违法', na=False)]
    
    all_counties = sorted(df['县级行政区名称'].unique().tolist())
    
    # 构建逐县数据
    county_data = []
    for county in all_counties:
        orig = original[original['县级行政区名称'] == county]
        veri = verified[verified['县级行政区名称'] == county]
        ill = illegal[illegal['县级行政区名称'] == county]
        grain = grain_illegal[grain_illegal['县级行政区名称'] == county]
        
        county_data.append({
            '县（区）': county,
            '原始下发个数': len(orig),
            '原始下发耕地面积': round(orig['耕地面积（亩）'].sum(), 2),
            '原始下发基本农田': round(orig['永久基本农田面积（亩）'].sum(), 2),
            '核实后个数': len(veri),
            '核实后耕地面积': round(veri['耕地面积（亩）'].sum(), 2),
            '核实后基本农田': round(veri['永久基本农田面积（亩）'].sum(), 2),
            '非农违法个数': len(ill),
            '非农违法耕地面积': round(ill['非农违法耕地面积（亩）'].sum(), 2),
            '非农违法基本农田': round(ill['非农违法永久基本农田面积（亩）'].sum(), 2),
            '非粮违法个数': len(grain),
            '非粮违法耕地面积': round(grain['非粮违法耕地面积（亩）'].sum(), 2),
            '非粮违法基本农田': round(grain['非粮违法永久基本农田面积（亩）'].sum(), 2),
        })
    
    # 按非农违法耕地面积降序排列
    county_data.sort(key=lambda x: x['非农违法耕地面积'], reverse=True)
    
    # 计算合计
    totals = {}
    for key in county_data[0].keys():
        if key == '县（区）':
            totals[key] = '合计'
        elif '个数' in key:
            totals[key] = sum(d[key] for d in county_data)
        else:
            totals[key] = round(sum(d[key] for d in county_data), 2)
    
    # 创建Excel
    wb = Workbook()
    ws = wb.active
    ws.title = '2026年7月大图斑逐县统计'
    
    # 样式定义
    header_font = Font(name='微软雅黑', size=10, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    data_font = Font(name='微软雅黑', size=9)
    total_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
    total_font = Font(name='微软雅黑', size=9, bold=True)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)
    
    # 标题
    ws.merge_cells('A1:M1')
    title_cell = ws['A1']
    title_cell.value = '2026年7月大图斑逐县统计表（按非农违法耕地面积降序）'
    title_cell.font = Font(name='微软雅黑', size=12, bold=True)
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 28
    
    # 表头第一行（大类）
    categories = {1: '县（区）', 2: '原始下发', 5: '核实后', 8: '非农违法', 11: '非粮违法'}
    for col, val in categories.items():
        cell = ws.cell(row=2, column=col, value=val)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    
    ws.merge_cells('A2:A3')
    ws.merge_cells('B2:D2')
    ws.merge_cells('E2:G2')
    ws.merge_cells('H2:J2')
    ws.merge_cells('K2:M2')
    
    # 表头第二行（子项）
    sub_headers = {
        2: '个数', 3: '耕地面积（亩）', 4: '基本农田（亩）',
        5: '个数', 6: '耕地面积（亩）', 7: '基本农田（亩）',
        8: '个数', 9: '耕地面积（亩）', 10: '基本农田（亩）',
        11: '个数', 12: '耕地面积（亩）', 13: '基本农田（亩）'
    }
    for col, val in sub_headers.items():
        cell = ws.cell(row=3, column=col, value=val)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    
    # 列宽
    col_widths = {'A': 11, 'B': 6, 'C': 12, 'D': 12, 'E': 6, 'F': 12, 'G': 12, 
                  'H': 6, 'I': 14, 'J': 14, 'K': 6, 'L': 14, 'M': 14}
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width
    
    # 数据行
    row_num = 4
    for d in county_data:
        row_data = [
            d['县（区）'],
            d['原始下发个数'], d['原始下发耕地面积'], d['原始下发基本农田'],
            d['核实后个数'], d['核实后耕地面积'], d['核实后基本农田'],
            d['非农违法个数'], d['非农违法耕地面积'], d['非农违法基本农田'],
            d['非粮违法个数'], d['非粮违法耕地面积'], d['非粮违法基本农田']
        ]
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = left_align if col_idx == 1 else center_align
        row_num += 1
    
    # 合计行
    total_row = row_num
    total_row_data = [
        totals['县（区）'],
        totals['原始下发个数'], totals['原始下发耕地面积'], totals['原始下发基本农田'],
        totals['核实后个数'], totals['核实后耕地面积'], totals['核实后基本农田'],
        totals['非农违法个数'], totals['非农违法耕地面积'], totals['非农违法基本农田'],
        totals['非粮违法个数'], totals['非粮违法耕地面积'], totals['非粮违法基本农田']
    ]
    for col_idx, val in enumerate(total_row_data, 1):
        cell = ws.cell(row=total_row, column=col_idx, value=val)
        cell.font = total_font
        cell.fill = total_fill
        cell.border = thin_border
        cell.alignment = left_align if col_idx == 1 else center_align
    
    ws.row_dimensions[total_row].height = 22
    ws.freeze_panes = 'B4'
    
    # A4横向打印设置
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4
    ws.page_margins.header = 0.3
    ws.page_margins.footer = 0.3
    
    # 保存
    wb.save(output_file)
    print(f'Excel文件已生成：{output_file}')
    print(f'共 {len(county_data)} 个县（区）+ 1行合计')
    print(f'按非农违法耕地面积降序排列')


if __name__ == '__main__':
    if len(sys.argv) < 3:
        print('用法：python generate_county_stats.py <输入文件> <输出文件>')
        print('示例：python generate_county_stats.py "遥感检测检索明细表.xlsx" "逐县统计表.xlsx"')
        sys.exit(1)
    
    input_file = sys.argv[1]
    output_file = sys.argv[2]
    
    if not os.path.exists(input_file):
        print(f'错误：输入文件不存在：{input_file}')
        sys.exit(1)
    
    generate_county_stats(input_file, output_file)
