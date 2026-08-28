# -*- coding: utf-8 -*-
"""
遥感监测违法线索统计脚本
生成非农违法/非粮违法的整改统计表，包含校核列
"""

import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
import sys
import os

def generate_illegal_stats(input_file, output_file, illegal_type='非农违法'):
    """
    生成违法线索统计表
    
    参数:
        input_file: 输入的遥感检测检索明细表Excel文件路径
        output_file: 输出的统计表Excel文件路径
        illegal_type: 违法类型，'非农违法' 或 '非粮违法'
    """
    
    # 读取数据
    df = pd.read_excel(input_file)
    
    # 数据筛选
    verified = df[df['图斑状态'].isna() | (df['图斑状态'] == '拆分')]
    
    if illegal_type == '非农违法':
        illegal = verified[verified['核实认定意见'].str.contains('非农违法', na=False)]
        area_col = '非农违法面积（亩）'
        farmland_col = '非农违法耕地面积（亩）'
        basic_farmland_col = '非农违法永久基本农田面积（亩）'
    else:
        illegal = verified[verified['核实认定意见'].str.contains('非粮违法', na=False)]
        area_col = '非粮违法面积（亩）'
        farmland_col = '非粮违法耕地面积（亩）'
        basic_farmland_col = '非粮违法永久基本农田面积（亩）'
    
    # 整改判定
    completed = illegal[(illegal['剩余未整改面积'] == 0) & (illegal['拆分状态'] == '审核结束')]
    demolish = completed[completed['整改情况'].str.contains('拆除', na=False) | (completed['拆除面积（亩）'] > 0)]
    supplement = completed[completed['整改情况'].str.contains('补办', na=False) | (completed['补办面积（亩）'] > 0)]
    pending = illegal[(illegal['剩余未整改面积'] > 0) | (illegal['拆分状态'] != '审核结束')]
    
    all_counties = sorted(illegal['县级行政区名称'].unique().tolist())
    
    # 构建逐县数据
    county_data = []
    for county in all_counties:
        ill = illegal[illegal['县级行政区名称'] == county]
        comp = completed[completed['县级行政区名称'] == county]
        dem = demolish[demolish['县级行政区名称'] == county]
        sup = supplement[supplement['县级行政区名称'] == county]
        pend = pending[pending['县级行政区名称'] == county]
        
        # 计算各项指标
        total_count = len(ill)
        total_area = round(ill[area_col].sum(), 2)
        total_farmland = round(ill[farmland_col].sum(), 2)
        
        completed_count = len(comp)
        completed_area = round(comp[area_col].sum(), 2)
        completed_farmland = round(comp[farmland_col].sum(), 2)
        
        demolish_count = len(dem)
        demolish_area = round(dem['拆除面积（亩）'].sum(), 2)
        demolish_farmland = round(dem['拆除耕地面积（亩）'].sum(), 2)
        
        supplement_count = len(sup)
        supplement_area = round(sup['补办面积（亩）'].sum(), 2)
        supplement_farmland = round(sup['补办耕地面积（亩）'].sum(), 2)
        
        pending_count = len(pend)
        pending_area = round(pend[area_col].sum(), 2)
        pending_farmland = round(pend[farmland_col].sum(), 2)
        
        # 校核
        check1_count = total_count - (completed_count + pending_count)
        check1_area = total_area - (completed_area + pending_area)
        check1_farmland = total_farmland - (completed_farmland + pending_farmland)
        
        check2_count = completed_count - (demolish_count + supplement_count)
        check2_area = completed_area - (demolish_area + supplement_area)
        check2_farmland = completed_farmland - (demolish_farmland + supplement_farmland)
        
        county_data.append({
            '县（区）': county,
            '总任务个数': total_count,
            '总任务面积': total_area,
            '总任务耕地面积': total_farmland,
            '已整改个数': completed_count,
            '已整改面积': completed_area,
            '已整改耕地面积': completed_farmland,
            '拆除个数': demolish_count,
            '拆除面积': demolish_area,
            '拆除耕地面积': demolish_farmland,
            '补办个数': supplement_count,
            '补办面积': supplement_area,
            '补办耕地面积': supplement_farmland,
            '未整改个数': pending_count,
            '未整改面积': pending_area,
            '未整改耕地面积': pending_farmland,
            '校核1_个数差异': check1_count,
            '校核1_面积差异': check1_area,
            '校核1_耕地差异': check1_farmland,
            '校核2_个数差异': check2_count,
            '校核2_面积差异': check2_area,
            '校核2_耕地差异': check2_farmland,
        })
    
    # 按未整改耕地面积降序排列
    county_data.sort(key=lambda x: x['未整改耕地面积'], reverse=True)
    
    # 计算合计
    totals = {}
    for key in county_data[0].keys():
        if key == '县（区）':
            totals[key] = '总计'
        elif '个数' in key or '差异' in key:
            totals[key] = sum(d[key] for d in county_data)
        else:
            totals[key] = round(sum(d[key] for d in county_data), 2)
    
    # 创建Excel
    wb = Workbook()
    ws = wb.active
    ws.title = f'2026年遥感监测违法线索统计表（{illegal_type}）'
    
    # 样式定义
    header_font = Font(name='微软雅黑', size=9, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    sub_header_fill = PatternFill(start_color='D6E4F0', end_color='D6E4F0', fill_type='solid')
    sub_header_font = Font(name='微软雅黑', size=8, bold=True)
    data_font = Font(name='微软雅黑', size=8)
    total_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
    total_font = Font(name='微软雅黑', size=8, bold=True)
    error_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
    error_font = Font(name='微软雅黑', size=8, color='9C0006')
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)
    
    # 标题
    ws.merge_cells('A1:T1')
    title_cell = ws['A1']
    title_cell.value = f'2026年遥感监测违法线索统计表（{illegal_type}）'
    title_cell.font = Font(name='微软雅黑', size=12, bold=True)
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 28
    
    # 单位说明
    unit_cell = ws.cell(row=1, column=21)
    unit_cell.value = '单位：个、亩'
    unit_cell.font = Font(name='微软雅黑', size=8)
    
    # 表头第一行（大类）
    categories = {
        1: '序号', 2: '县（市、区）',
        3: '总任务', 6: '已整改', 9: '拆除', 12: '补办', 15: '未整改', 18: '校核'
    }
    for col, val in categories.items():
        cell = ws.cell(row=2, column=col, value=val)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    
    ws.merge_cells('A2:A3')
    ws.merge_cells('B2:B3')
    ws.merge_cells('C2:E2')
    ws.merge_cells('F2:H2')
    ws.merge_cells('I2:K2')
    ws.merge_cells('L2:N2')
    ws.merge_cells('O2:Q2')
    ws.merge_cells('R2:U2')
    
    # 表头第二行（子项）
    sub_headers = {
        3: '个数', 4: '面积', 5: '耕地面积',
        6: '个数', 7: '面积', 8: '耕地面积',
        9: '个数', 10: '面积', 11: '耕地面积',
        12: '个数', 13: '面积', 14: '耕地面积',
        15: '个数', 16: '面积', 17: '耕地面积',
        18: '总=已+未\n个数', 19: '总=已+未\n面积', 20: '总=已+未\n耕地', 21: '已=拆+补\n个数'
    }
    for col, val in sub_headers.items():
        cell = ws.cell(row=3, column=col, value=val)
        cell.font = sub_header_font
        cell.fill = sub_header_fill
        cell.alignment = center_align
        cell.border = thin_border
    
    # 列宽
    col_widths = {
        'A': 4, 'B': 9,
        'C': 5, 'D': 7, 'E': 9,
        'F': 5, 'G': 7, 'H': 9,
        'I': 5, 'J': 7, 'K': 9,
        'L': 5, 'M': 7, 'N': 9,
        'O': 5, 'P': 7, 'Q': 9,
        'R': 7, 'S': 7, 'T': 7, 'U': 7
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width
    
    # 数据行
    row_num = 4
    for idx, d in enumerate(county_data, 1):
        row_data = [
            idx, d['县（区）'],
            d['总任务个数'], d['总任务面积'], d['总任务耕地面积'],
            d['已整改个数'], d['已整改面积'], d['已整改耕地面积'],
            d['拆除个数'], d['拆除面积'], d['拆除耕地面积'],
            d['补办个数'], d['补办面积'], d['补办耕地面积'],
            d['未整改个数'], d['未整改面积'], d['未整改耕地面积'],
            d['校核1_个数差异'], d['校核1_面积差异'], d['校核1_耕地差异'],
            d['校核2_个数差异']
        ]
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if col_idx in [1, 2]:
                cell.alignment = left_align if col_idx == 2 else center_align
            else:
                cell.alignment = center_align
            # 校核列标红
            if col_idx >= 18 and val != 0:
                cell.fill = error_fill
                cell.font = error_font
        row_num += 1
    
    # 合计行
    total_row = row_num
    total_row_data = [
        '', totals['县（区）'],
        totals['总任务个数'], totals['总任务面积'], totals['总任务耕地面积'],
        totals['已整改个数'], totals['已整改面积'], totals['已整改耕地面积'],
        totals['拆除个数'], totals['拆除面积'], totals['拆除耕地面积'],
        totals['补办个数'], totals['补办面积'], totals['补办耕地面积'],
        totals['未整改个数'], totals['未整改面积'], totals['未整改耕地面积'],
        totals['校核1_个数差异'], totals['校核1_面积差异'], totals['校核1_耕地差异'],
        totals['校核2_个数差异']
    ]
    for col_idx, val in enumerate(total_row_data, 1):
        cell = ws.cell(row=total_row, column=col_idx, value=val)
        cell.font = total_font
        cell.fill = total_fill
        cell.border = thin_border
        if col_idx == 2:
            cell.alignment = left_align
        else:
            cell.alignment = center_align
        # 校核列标红
        if col_idx >= 18 and val != 0:
            cell.fill = error_fill
            cell.font = error_font
    
    ws.row_dimensions[total_row].height = 22
    ws.freeze_panes = 'C4'
    
    # A4横向打印设置
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.3
    ws.page_margins.bottom = 0.3
    ws.page_margins.header = 0.2
    ws.page_margins.footer = 0.2
    
    # 保存
    wb.save(output_file)
    print(f'Excel文件已生成：{output_file}')
    print(f'共 {len(county_data)} 个县（区）+ 1行总计')
    print(f'按未整改耕地面积降序排列')
    print()
    print('=== 数据概览 ===')
    print(f'总任务：{totals["总任务个数"]}个，{totals["总任务面积"]}亩，耕地{totals["总任务耕地面积"]}亩')
    print(f'已整改：{totals["已整改个数"]}个，{totals["已整改面积"]}亩，耕地{totals["已整改耕地面积"]}亩')
    print(f'  - 拆除：{totals["拆除个数"]}个，{totals["拆除面积"]}亩，耕地{totals["拆除耕地面积"]}亩')
    print(f'  - 补办：{totals["补办个数"]}个，{totals["补办面积"]}亩，耕地{totals["补办耕地面积"]}亩')
    print(f'未整改：{totals["未整改个数"]}个，{totals["未整改面积"]}亩，耕地{totals["未整改耕地面积"]}亩')
    print()
    print('=== 校核结果 ===')
    print(f'校核1（总任务=已整改+未整改）：')
    print(f'  个数差异：{totals["校核1_个数差异"]}')
    print(f'  面积差异：{totals["校核1_面积差异"]}')
    print(f'  耕地差异：{totals["校核1_耕地差异"]}')
    print(f'校核2（已整改=拆除+补办）：')
    print(f'  个数差异：{totals["校核2_个数差异"]}')
    print(f'  面积差异：{totals["校核2_面积差异"]}')
    print(f'  耕地差异：{totals["校核2_耕地差异"]}')


if __name__ == '__main__':
    if len(sys.argv) < 3:
        print('用法：python generate_illegal_stats.py <输入文件> <输出文件> [违法类型]')
        print('示例：python generate_illegal_stats.py "遥感检测检索明细表.xlsx" "违法线索统计表.xlsx" "非农违法"')
        print('违法类型可选：非农违法（默认）、非粮违法')
        sys.exit(1)
    
    input_file = sys.argv[1]
    output_file = sys.argv[2]
    illegal_type = sys.argv[3] if len(sys.argv) > 3 else '非农违法'
    
    if illegal_type not in ['非农违法', '非粮违法']:
        print(f'错误：违法类型必须为"非农违法"或"非粮违法"，当前为：{illegal_type}')
        sys.exit(1)
    
    if not os.path.exists(input_file):
        print(f'错误：输入文件不存在：{input_file}')
        sys.exit(1)
    
    generate_illegal_stats(input_file, output_file, illegal_type)
