# -*- coding: utf-8 -*-
"""变更调查新增建设线索 · 各地类情况汇总

按「实地地类」汇总：图斑个数 / 个数占比 / 面积(亩) / 面积占比 /
平均单斑面积(亩) / 涉及县区数 / 主要分布县区(前3)

排序：图斑个数降序；「（未认定）」置末

用法：python generate_dilei_summary.py <源xlsx> [统计日期YYYY-MM-DD]
输出：<源目录>/变更调查各地类情况统计_<日期>.xlsx
"""
import sys
import os
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from dilei_map import DIL_EI_MAP

if len(sys.argv) < 2:
    print('用法: python generate_dilei_summary.py <源xlsx> [统计日期YYYY-MM-DD]')
    sys.exit(1)
SRC = sys.argv[1]
OUT_DIR = os.path.dirname(SRC)
TODAY = sys.argv[2] if len(sys.argv) > 2 else datetime.now().strftime('%Y-%m-%d')

UNDEF = '（未认定）'
COLS = ['序号', '实地地类', '图斑个数', '个数占比', '面积(亩)', '面积占比',
        '平均单斑面积(亩)', '涉及县区数', '主要分布县区(前3)']


def to_name(v):
    if pd.isna(v):
        return UNDEF
    return DIL_EI_MAP.get(str(v).strip(), f'{v} (?)')


def build_summary(df):
    total_cnt = len(df)
    total_area = float(df['__area'].sum())

    rows = []
    for dl, sub in df.groupby('__dilei'):
        cnt = len(sub)
        area = float(sub['__area'].sum())
        vc = sub['__county'].value_counts()
        rows.append({
            '实地地类': dl,
            '图斑个数': cnt,
            '个数占比': cnt / total_cnt if total_cnt else 0,
            '面积(亩)': round(area, 2),
            '面积占比': area / total_area if total_area else 0,
            '平均单斑面积(亩)': round(area / cnt, 2) if cnt else 0,
            '涉及县区数': int(vc.size),
            '主要分布县区(前3)': '、'.join(f'{k}({v})' for k, v in vc.head(3).items()),
        })

    out = pd.DataFrame(rows).sort_values('图斑个数', ascending=False).reset_index(drop=True)
    # 未认定置末
    if UNDEF in out['实地地类'].values:
        undef_row = out[out['实地地类'] == UNDEF]
        out = pd.concat([out[out['实地地类'] != UNDEF], undef_row], ignore_index=True)
    out.insert(0, '序号', range(1, len(out) + 1))

    total_row = {
        '序号': '', '实地地类': '合计', '图斑个数': total_cnt, '个数占比': 1.0,
        '面积(亩)': round(total_area, 2), '面积占比': 1.0, '平均单斑面积(亩)': '',
        '涉及县区数': int(df['__county'].nunique()),
        '主要分布县区(前3)': '',
    }
    return out, total_row


def write_sheet(ws, df, total_row, title):
    header_font = Font(name='微软雅黑', size=10, bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='4472C4')
    cell_font = Font(name='微软雅黑', size=9)
    total_font = Font(name='微软雅黑', size=9, bold=True)
    total_fill = PatternFill('solid', fgColor='FFF2CC')
    thin = Side(border_style='thin', color='B4B4B4')
    border = Border(top=thin, left=thin, right=thin, bottom=thin)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = Alignment(horizontal='left', vertical='center', wrap_text=True)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLS))
    ws.cell(1, 1, f'{title}（{TODAY}）').font = Font(name='微软雅黑', size=14, bold=True)
    ws.cell(1, 1).alignment = center
    ws.row_dimensions[1].height = 26

    for i, c in enumerate(COLS, start=1):
        cell = ws.cell(2, i, c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border
    ws.row_dimensions[2].height = 28

    r = 3
    for row in df.to_dict('records'):
        for i, c in enumerate(COLS, start=1):
            v = row[c]
            cell = ws.cell(r, i, v)
            cell.font = cell_font
            cell.border = border
            cell.alignment = left if c in ('实地地类', '主要分布县区(前3)') else center
            if c in ('个数占比', '面积占比'):
                cell.number_format = '0.0%'
            elif c in ('面积(亩)', '平均单斑面积(亩)'):
                cell.number_format = '0.00'
        ws.row_dimensions[r].height = 20
        r += 1

    for i, c in enumerate(COLS, start=1):
        v = total_row[c]
        cell = ws.cell(r, i, v)
        cell.font = total_font
        cell.fill = total_fill
        cell.border = border
        cell.alignment = left if c in ('实地地类', '主要分布县区(前3)') else center
        if c in ('个数占比', '面积占比'):
            cell.number_format = '0.0%'
        elif c == '面积(亩)':
            cell.number_format = '0.00'
    ws.row_dimensions[r].height = 20

    widths = [6, 16, 10, 10, 12, 10, 14, 11, 46]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = 'C3'
    ws.auto_filter.ref = f'A2:{get_column_letter(len(COLS))}{r - 1}'

    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4


def main():
    df = pd.read_excel(SRC)
    print(f'源表行数: {len(df)}')

    df['__county'] = df['行政区名称'].fillna(df['县级行政区名称']).fillna('（未分配）')
    df['__dilei'] = df['实地地类'].apply(to_name)
    df['__area'] = pd.to_numeric(df['图斑面积(亩)'], errors='coerce').fillna(0)

    out, total_row = build_summary(df)

    print('\n--- 各地类情况 ---')
    for row in out.to_dict('records'):
        print(f"  {row['实地地类']}: {row['图斑个数']} 个 / {row['面积(亩)']} 亩 / "
              f"{row['个数占比']*100:.1f}% / 涉及 {row['涉及县区数']} 县区")

    wb = Workbook()
    ws = wb.active
    ws.title = '各地类情况'
    write_sheet(ws, out, total_row, '变更调查新增建设线索·各地类情况统计')

    out_path = os.path.join(OUT_DIR, f'变更调查各地类情况统计_{TODAY}.xlsx')
    wb.save(out_path)
    print(f'\n输出: {out_path}')


if __name__ == '__main__':
    main()
