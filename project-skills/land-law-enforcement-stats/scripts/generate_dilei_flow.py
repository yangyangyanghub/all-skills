# -*- coding: utf-8 -*-
"""变更调查新增建设线索 · 县 × 实地地类 流向统计

输出 2 个 Sheet：
- Sheet1 图斑个数：县名称 × 地类 交叉计数，含合计行/合计列
- Sheet2 图斑面积：县名称 × 地类 面积合计（亩，用 图斑面积(亩) 全量列）

排序：地类列按总计降序，县行按总计降序

用法：python generate_dilei_flow.py <源xlsx> [统计日期YYYY-MM-DD]
输出：<源目录>/变更调查地类流向统计_<日期>.xlsx
"""
import sys
import os
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from dilei_map import DIL_EI_MAP

if len(sys.argv) < 2:
    print('用法: python generate_dilei_flow.py <源xlsx> [统计日期YYYY-MM-DD]')
    sys.exit(1)
SRC = sys.argv[1]
OUT_DIR = os.path.dirname(SRC)
TODAY = sys.argv[2] if len(sys.argv) > 2 else datetime.now().strftime('%Y-%m-%d')


def main():
    df = pd.read_excel(SRC)
    print(f'源表行数: {len(df)}')

    # 县分组（行政区名称，979 全非空）
    df['__county'] = df['行政区名称'].fillna(df['县级行政区名称']).fillna('（未分配）')

    # 地类映射（30 个 NaN → 未认定）
    def to_name(v):
        if pd.isna(v):
            return '（未认定）'
        return DIL_EI_MAP.get(str(v).strip(), f'{v} (?)')

    df['__dilei'] = df['实地地类'].apply(to_name)

    # 面积用 图斑面积(亩)（979 全非空）
    df['__area'] = pd.to_numeric(df['图斑面积(亩)'], errors='coerce').fillna(0)

    # ---- 透视：个数 ----
    pivot_cnt = df.pivot_table(index='__county', columns='__dilei', values='线索编号',
                               aggfunc='count', fill_value=0)
    # ---- 透视：面积 ----
    pivot_area = df.pivot_table(index='__county', columns='__dilei', values='__area',
                                aggfunc='sum', fill_value=0)

    # 地类列按总计降序（两个表列顺序一致）
    col_order = pivot_cnt.sum(axis=0).sort_values(ascending=False).index.tolist()
    pivot_cnt = pivot_cnt[col_order]
    pivot_area = pivot_area[col_order]

    # 县行按总计降序
    row_order = pivot_cnt.sum(axis=1).sort_values(ascending=False).index.tolist()
    pivot_cnt = pivot_cnt.loc[row_order]
    pivot_area = pivot_area.loc[row_order]

    # 加合计行/列
    def add_totals(p):
        p = p.copy()
        p.loc['合计'] = p.sum(axis=0)
        p['合计'] = p.sum(axis=1)
        return p

    pivot_cnt = add_totals(pivot_cnt)
    pivot_area = add_totals(pivot_area)

    print('\n--- 各地类总计（个数）---')
    for c in pivot_cnt.columns[:-1]:
        print(f'  {c}: {int(pivot_cnt.loc["合计", c])} 个 / {pivot_area.loc["合计", c]:.2f} 亩')

    # ---- 写 Excel ----
    wb = Workbook()
    header_font = Font(name='微软雅黑', size=10, bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='4472C4')
    cell_font = Font(name='微软雅黑', size=9)
    total_font = Font(name='微软雅黑', size=9, bold=True)
    total_fill = PatternFill('solid', fgColor='FFF2CC')
    thin = Side(border_style='thin', color='B4B4B4')
    border = Border(top=thin, left=thin, right=thin, bottom=thin)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = Alignment(horizontal='left', vertical='center', wrap_text=True)

    # Sheet1 个数
    ws1 = wb.active
    ws1.title = '图斑个数'
    write_sheet(ws1, pivot_cnt, '变更调查新增建设线索·县区地类流向统计（图斑个数）', TODAY, center)
    # Sheet2 面积
    ws2 = wb.create_sheet('图斑面积')
    write_sheet(ws2, pivot_area, '变更调查新增建设线索·县区地类流向统计（面积·亩）', TODAY, center)

    for ws in (ws1, ws2):
        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = 0.3
        ws.page_margins.right = 0.3
        ws.page_margins.top = 0.4
        ws.page_margins.bottom = 0.4

    out_path = os.path.join(OUT_DIR, f'变更调查地类流向统计_{TODAY}.xlsx')
    wb.save(out_path)
    print(f'\n输出: {out_path}')


def write_sheet(ws, pivot, title, today, center):
    """写入交叉表：标题行 + 表头 + 数据行（含合计行/列）"""
    n_dilei = len(pivot.columns) - 1          # 地类列数（不含合计列）
    n_county = len(pivot.index) - 1           # 县行数（不含合计行）
    total_cols = n_dilei + 2                  # 县名称 + 地类列 + 合计列

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    ws.cell(1, 1, f'{title}（{today}）').font = Font(name='微软雅黑', size=14, bold=True)
    ws.cell(1, 1).alignment = center
    ws.row_dimensions[1].height = 26

    # 表头（第2行）
    header_font = Font(name='微软雅黑', size=9, bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='4472C4')
    border = Border(top=Side(border_style='thin', color='B4B4B4'),
                    left=Side(border_style='thin', color='B4B4B4'),
                    right=Side(border_style='thin', color='B4B4B4'),
                    bottom=Side(border_style='thin', color='B4B4B4'))
    cell_font = Font(name='微软雅黑', size=9)
    total_font = Font(name='微软雅黑', size=9, bold=True)
    total_fill = PatternFill('solid', fgColor='FFF2CC')
    left = Alignment(horizontal='left', vertical='center', wrap_text=True)

    ws.cell(2, 1, '县（区）')
    for i, c in enumerate(pivot.columns):
        col = i + 2
        ws.cell(2, col, str(c))
        ws.cell(2, col).font = header_font
        ws.cell(2, col).fill = header_fill
        ws.cell(2, col).alignment = center
        ws.cell(2, col).border = border
    ws.cell(2, 1).font = header_font
    ws.cell(2, 1).fill = header_fill
    ws.cell(2, 1).alignment = center
    ws.cell(2, 1).border = border
    ws.row_dimensions[2].height = 28

    # 数据行
    for r, (county, row) in enumerate(pivot.iterrows(), start=3):
        is_total = (county == '合计')
        ws.cell(r, 1, county)
        for i, c in enumerate(pivot.columns):
            v = row[c]
            col = i + 2
            ws.cell(r, col, int(v) if is_int_like(v) else round(float(v), 2))
            ws.cell(r, col).alignment = center
            ws.cell(r, col).border = border
            ws.cell(r, col).font = total_font if is_total else cell_font
        ws.cell(r, 1).alignment = left
        ws.cell(r, 1).border = border
        ws.cell(r, 1).font = total_font if is_total else cell_font
        if is_total:
            for col in range(1, total_cols + 1):
                ws.cell(r, col).fill = total_fill
        ws.row_dimensions[r].height = 20

    # 列宽
    ws.column_dimensions['A'].width = 14
    for i in range(1, n_dilei + 1):
        ws.column_dimensions[get_column_letter(i + 1)].width = 11
    ws.column_dimensions[get_column_letter(total_cols)].width = 11

    ws.freeze_panes = 'B3'
    ws.auto_filter.ref = f'A2:{get_column_letter(total_cols)}{2 + n_county}'


def is_int_like(v):
    return isinstance(v, (int, float)) and float(v).is_integer() and not isinstance(v, bool)


if __name__ == '__main__':
    main()
