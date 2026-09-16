# 面积核对：核查填报（图斑列表导出） vs 套合结果（GIS 真值）
#
# 校验两条：
#   1. 逐项比对：填报值 vs 套合值，差异 > 容差 → 标差异
#   2. 公式校验：地块面积 = 耕地 + 非耕农用地 + 未利用地 + 建设用地
#      （公式来源：套合表实测 1548/1548 匹配；源自《一体化系统填报拆分流程》DKMJ=GDMJ+FGNYDMJ+WLYDMJ+建设用地）
#
# ⚠️ 注意：导出表列名用全角括号（），套合表用半角括号()，不要混用
import argparse
import json
import os
import sys


def _num(v):
  if v is None or v == "":
    return None
  try:
    return float(v)
  except (TypeError, ValueError):
    return None


def _norm(v):
  return 0.0 if v is None else v


def load_sheet(path, sheet_name):
  import openpyxl
  wb = openpyxl.load_workbook(path, data_only=True)
  ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.worksheets[0]
  it = ws.iter_rows(values_only=True)
  header = [str(c).strip() if c is not None else "" for c in next(it)]
  rows = [dict(zip(header, r)) for r in it]
  wb.close()
  return rows


# 核查填报（图斑列表导出）→ 内部字段名
FILLED_MAP = {
  "areaMu": "地块面积",
  "cultivated": "占耕地面积",
  "basicFarmland": "占基本农田面积",
  "nonCultivatedAgri": "占非耕农用地面积",
  "unused": "占未利用地面积",
  "ecoRedline": "占生态红线面积",
  "construction": "建设用地面积（亩）",
}

# 套合结果表 → 内部字段名
TRUTH_MAP = {
  "areaMu": "地块总面积(亩)",
  "cultivated": "耕地面积(亩)",
  "basicFarmland": "基本农田面积(亩)",
  "nonCultivatedAgri": "非耕农用地面积(亩)",
  "unused": "未利用地面积(亩)",
  "ecoRedline": "生态红线面积(亩)",
  "construction": "建设用地面积(亩)",
}

FIELD_LABEL = {
  "areaMu": "地块面积",
  "cultivated": "占耕地面积",
  "basicFarmland": "占基本农田面积",
  "nonCultivatedAgri": "占非耕农用地面积",
  "unused": "占未利用地面积",
  "ecoRedline": "占生态红线面积",
  "construction": "建设用地面积",
}


def build_index(rows, key_col, col_map):
  out = {}
  for r in rows:
    k = r.get(key_col)
    if k is None or str(k).strip() == "":
      continue
    out[str(k).strip()] = {f: _num(r.get(c)) for f, c in col_map.items()}
  return out


def verify(filled_path, truth_path, tolerance=0.01,
           filled_sheet="图斑列表", truth_sheet="套合面积审核表",
           filled_key="线索编号", truth_key="图斑标识"):
  filled = build_index(load_sheet(filled_path, filled_sheet), filled_key, FILLED_MAP)
  truth = build_index(load_sheet(truth_path, truth_sheet), truth_key, TRUTH_MAP)
  both = sorted(set(filled) & set(truth))

  field_diffs = []
  formula_fails = []
  formula_checked = 0
  formula_skipped = 0
  for k in both:
    f, t = filled[k], truth[k]
    # 1. 逐项比对
    for field in FILLED_MAP:
      a, b = f.get(field), t.get(field)
      if a is None or b is None:
        continue
      if abs(a - b) > tolerance:
        field_diffs.append({
          "id": k, "field": FIELD_LABEL[field],
          "filled": round(a, 4), "truth": round(b, 4),
          "diff": round(a - b, 4),
        })
    # 2. 公式校验（用填报值；需 5 字段齐全）
    dk, gd, fg, wy, js = (f.get("areaMu"), f.get("cultivated"),
                          f.get("nonCultivatedAgri"), f.get("unused"), f.get("construction"))
    if None in (dk, gd, fg, wy, js):
      formula_skipped += 1
      continue
    formula_checked += 1
    s = _norm(gd) + _norm(fg) + _norm(wy) + _norm(js)
    if abs(dk - s) > tolerance:
      formula_fails.append({
        "id": k, "areaMu": round(dk, 4), "sum": round(s, 4),
        "cultivated": round(_norm(gd), 4), "nonCultivatedAgri": round(_norm(fg), 4),
        "unused": round(_norm(wy), 4), "construction": round(_norm(js), 4),
        "diff": round(dk - s, 4),
      })

  return {
    "filledCount": len(filled),
    "truthCount": len(truth),
    "comparableCount": len(both),
    "tolerance": tolerance,
    "fieldDiffCount": len(field_diffs),
    "fieldDiffs": field_diffs,
    "formulaCheckedCount": formula_checked,
    "formulaSkippedCount": formula_skipped,
    "formulaFailCount": len(formula_fails),
    "formulaFails": formula_fails,
  }


def render(result):
  lines = ["## 面积核对结果", ""]
  lines.append(f"- 核查填报表：**{result['filledCount']}** 个图斑")
  lines.append(f"- 套合结果表：**{result['truthCount']}** 个图斑")
  lines.append(f"- 可比对：**{result['comparableCount']}** 个")
  lines.append(f"- 容差：**{result['tolerance']}** 亩")
  lines.append(f"- 逐项差异：**{result['fieldDiffCount']}** 处")
  fc = result.get("formulaCheckedCount", 0)
  fs = result.get("formulaSkippedCount", 0)
  lines.append(f"- 公式校验：可执行 **{fc}** 个 / 跳过 **{fs}** 个")
  if fs and fc == 0:
    lines.append("  - ⚠️ **全部跳过**——导出表缺「建设用地面积（亩）」（0% 填报），"
                 "无法离线做公式校验；如需校验请从系统读取该字段")
  lines.append(f"- 公式不通过：**{result['formulaFailCount']}** 个")
  lines.append("")

  if result["fieldDiffs"]:
    import collections
    by_field = collections.Counter(d["field"] for d in result["fieldDiffs"])
    lines.append("### 差异按字段分布")
    lines.append("")
    lines.append("| 字段 | 差异数 |")
    lines.append("|---|---|")
    for f, n in by_field.most_common():
      lines.append(f"| {f} | {n} |")
    lines.append("")
    lines.append("### 差异明细（前 30 条）")
    lines.append("")
    lines.append("| 线索编号 | 字段 | 填报 | 套合 | 差异 |")
    lines.append("|---|---|---|---|---|")
    for d in result["fieldDiffs"][:30]:
      lines.append(f"| {d['id']} | {d['field']} | {d['filled']} | {d['truth']} | {d['diff']} |")
    lines.append("")

  if result["formulaFails"]:
    lines.append("### 公式不通过（前 30 条）")
    lines.append("")
    lines.append("| 线索编号 | 地块面积 | 耕地 | 非耕农用地 | 未利用地 | 建设用地 | 分项和 | 差异 |")
    lines.append("|---|---|---|---|---|---|---|---|")
    for d in result["formulaFails"][:30]:
      lines.append(f"| {d['id']} | {d['areaMu']} | {d['cultivated']} | {d['nonCultivatedAgri']} "
                   f"| {d['unused']} | {d['construction']} | {d['sum']} | {d['diff']} |")
  return "\n".join(lines)


def main():
  ap = argparse.ArgumentParser(description="面积核对：核查填报 vs 套合结果")
  ap.add_argument("--filled", required=True, help="图斑列表导出表（含核查填报面积）")
  ap.add_argument("--truth", required=True, help="套合结果表（GIS 真值）")
  ap.add_argument("--tolerance", type=float, default=0.01, help="容差（亩），默认 0.01")
  ap.add_argument("--filled-sheet", default="图斑列表")
  ap.add_argument("--truth-sheet", default="套合面积审核表")
  ap.add_argument("--out", help="输出 JSON 路径")
  args = ap.parse_args()

  result = verify(args.filled, args.truth, args.tolerance,
                  args.filled_sheet, args.truth_sheet)
  print(render(result))
  if args.out:
    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    with open(args.out, "w", encoding="utf-8") as f:
      json.dump(result, f, ensure_ascii=False, indent=2)
    print(f"\n已写出 → {args.out}")


if __name__ == "__main__":
  main()
