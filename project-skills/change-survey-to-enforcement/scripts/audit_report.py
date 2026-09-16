# 审核意见报告：合并「规则判定 + 举证核查 + 面积核对」→ Excel
#
# 输出三个 Sheet：
#   Sheet1「审核意见」  图斑级：判定 / 一致性 / 举证情况 / 缺失举证 / 面积差异 / 审核建议
#   Sheet2「面积差异」  字段级：图斑 / 字段 / 填报 / 套合 / 差异 / 级别
#   Sheet3「缺失举证」  材料级：图斑 / 判定类别 / 缺失材料 / 补交要求
#
# 举证状态来源：
#   - 离线：无系统举证数据 → 只列"规则要求"，标注"待核（需读系统）"
#   - 在线：传 --evidence <json> → 对比后列出真正缺失的材料
import argparse
import json
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from judge import load_rules, judge_parcel  # noqa: E402
from load_parcels import load_from_xlsx  # noqa: E402
from verify_areas import verify  # noqa: E402

# 样式（沿用土地执法技能约定）
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10)
DANGER_FILL = PatternFill("solid", fgColor="FFC7CE")
WARN_FILL = PatternFill("solid", fgColor="FFEB9C")
OK_FILL = PatternFill("solid", fgColor="C6EFCE")

# 审核建议分级
SUGGEST_PASS = "通过"
SUGGEST_SUPPLEMENT = "退回补材料"
SUGGEST_VERIFY_AREA = "退回核实面积"
SUGGEST_REVIEW = "人工复核"


def _area_level(diff):
  d = abs(diff)
  if d > 1:
    return "严重"
  if d > 0.1:
    return "警告"
  return "提醒"


def _suggest(consistent, missing_evidence, area_diffs, suspicious):
  reasons = []
  if suspicious:
    reasons.append(SUGGEST_REVIEW)
  if not consistent:
    reasons.append(SUGGEST_REVIEW)
  if missing_evidence:
    reasons.append(SUGGEST_SUPPLEMENT)
  if area_diffs:
    reasons.append(SUGGEST_VERIFY_AREA)
  if not reasons:
    return SUGGEST_PASS
  # 优先级：人工复核 > 退回补材料 > 退回核实面积
  for s in (SUGGEST_REVIEW, SUGGEST_SUPPLEMENT, SUGGEST_VERIFY_AREA):
    if s in reasons:
      return s
  return SUGGEST_PASS


def build_audit(filled_xlsx, truth_xlsx, rules_path, evidence=None, tolerance=0.01):
  rules = load_rules(rules_path)
  parcels = load_from_xlsx(filled_xlsx)
  # 在线：把系统举证注入图斑对象，供 judge.py 的举证核查使用
  if evidence:
    for p in parcels:
      e = evidence.get(p["id"])
      if not e:
        continue
      atts = list(e.get("attachments") or [])
      # 批文附件也算附件举证
      atts.extend(e.get("批文") or [])
      p["evidence"] = {
        "photos": e.get("photos", 0),
        "attachments": atts,
        "overlay": e.get("overlay") or {},
      }
  judgments = {j["id"]: j for j in (judge_parcel(p, rules) for p in parcels)}
  area = verify(filled_xlsx, truth_xlsx, tolerance)

  # 面积差异按图斑聚合
  area_by_id = {}
  for d in area["fieldDiffs"]:
    area_by_id.setdefault(d["id"], []).append(d)

  rows = []
  missing_rows = []
  for pid, j in judgments.items():
    ev = evidence.get(pid) if evidence else None
    required = j["evidence"].get("required", [])
    have = j["evidence"].get("have", [])
    if ev is None:
      # 离线：不知道系统里有没有 → 只列要求，标"待核"
      missing = []
      pending = required
    else:
      missing = j["evidence"].get("missing", [])
      pending = []

    ad = area_by_id.get(pid, [])
    rows.append({
      "id": pid,
      "expected": j["expected"]["top"] + ("·" + j["expected"]["category"] if j["expected"].get("category") else ""),
      "filled": j["filled"]["top"] + ("·" + j["filled"]["category"] if j["filled"].get("category") else ""),
      "consistent": j["consistent"],
      "suspicious": j["suspicious"],
      "required": required,
      "have": have,
      "missing": missing,
      "pending": pending,
      "areaDiffs": ad,
      "suggest": _suggest(j["consistent"], missing, ad, j["suspicious"]),
      "note": j.get("note", ""),
      "basis": j.get("basis", ""),
    })
    for m in missing:
      missing_rows.append({"id": pid, "category": rows[-1]["expected"], "material": m,
                           "status": "缺失", "how": "补交后重新提报"})
    for m in pending:
      missing_rows.append({"id": pid, "category": rows[-1]["expected"], "material": m,
                           "status": "待核", "how": "需读系统确认是否已提供"})

  return {"rows": rows, "areaDiffs": area["fieldDiffs"], "missingRows": missing_rows,
          "areaSummary": {k: area[k] for k in
                          ("filledCount", "truthCount", "comparableCount", "tolerance",
                           "fieldDiffCount", "formulaCheckedCount", "formulaSkippedCount",
                           "formulaFailCount")}}


def _style_header(ws, ncol):
  for c in range(1, ncol + 1):
    cell = ws.cell(row=1, column=c)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
  ws.freeze_panes = "A2"


def write_excel(result, out_path):
  os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
  wb = openpyxl.Workbook()

  # ---- Sheet1 审核意见 ----
  ws1 = wb.active
  ws1.title = "审核意见"
  h1 = ["线索编号", "规则判定", "县级填报", "一致性", "存疑", "举证要求",
        "已提供", "⚠️缺失举证", "面积差异", "审核建议", "判定依据"]
  ws1.append(h1)
  for r in result["rows"]:
    area_txt = "；".join(f"{d['field']}:{d['filled']}≠{d['truth']}" for d in r["areaDiffs"])
    miss_txt = "；".join(r["missing"]) if r["missing"] else ("；".join(r["pending"]) + "（待核）" if r["pending"] else "")
    cons = "一致" if r["consistent"] else "不一致"
    ws1.append([r["id"], r["expected"], r["filled"], cons,
                "是" if r["suspicious"] else "",
                "；".join(r["required"]), "；".join(r["have"]), miss_txt,
                area_txt, r["suggest"], r["basis"]])
    row = ws1.max_row
    if r["missing"]:
      ws1.cell(row=row, column=8).fill = DANGER_FILL
    elif r["pending"]:
      ws1.cell(row=row, column=8).fill = WARN_FILL
    if not r["consistent"]:
      ws1.cell(row=row, column=4).fill = DANGER_FILL
    if r["areaDiffs"]:
      ws1.cell(row=row, column=9).fill = WARN_FILL
    if r["suggest"] == SUGGEST_PASS:
      ws1.cell(row=row, column=10).fill = OK_FILL
  _style_header(ws1, len(h1))
  for i, w in enumerate([20, 20, 20, 8, 6, 40, 24, 34, 30, 14, 40], 1):
    ws1.column_dimensions[get_column_letter(i)].width = w

  # ---- Sheet2 面积差异 ----
  ws2 = wb.create_sheet("面积差异")
  h2 = ["线索编号", "字段", "填报值", "套合值", "差异", "级别"]
  ws2.append(h2)
  for d in sorted(result["areaDiffs"], key=lambda x: -abs(x["diff"])):
    ws2.append([d["id"], d["field"], d["filled"], d["truth"], d["diff"], _area_level(d["diff"])])
    lvl = _area_level(d["diff"])
    fill = DANGER_FILL if lvl == "严重" else (WARN_FILL if lvl == "警告" else None)
    if fill:
      ws2.cell(row=ws2.max_row, column=6).fill = fill
  _style_header(ws2, len(h2))
  for i, w in enumerate([22, 18, 12, 12, 12, 10], 1):
    ws2.column_dimensions[get_column_letter(i)].width = w

  # ---- Sheet3 缺失举证 ----
  ws3 = wb.create_sheet("缺失举证")
  h3 = ["线索编号", "判定类别", "举证材料", "状态", "补交要求"]
  ws3.append(h3)
  for m in result["missingRows"]:
    ws3.append([m["id"], m["category"], m["material"], m["status"], m["how"]])
    if m["status"] == "缺失":
      ws3.cell(row=ws3.max_row, column=4).fill = DANGER_FILL
    else:
      ws3.cell(row=ws3.max_row, column=4).fill = WARN_FILL
  _style_header(ws3, len(h3))
  for i, w in enumerate([22, 22, 50, 10, 24], 1):
    ws3.column_dimensions[get_column_letter(i)].width = w

  wb.save(out_path)
  return out_path

def render_summary(result):
  rows = result["rows"]
  s = result["areaSummary"]
  lines = ["## 审核意见汇总", ""]
  lines.append(f"- 图斑总数：**{len(rows)}**")
  lines.append(f"- 与填报不一致：**{sum(1 for r in rows if not r['consistent'])}**")
  lines.append(f"- 存疑待人工：**{sum(1 for r in rows if r['suspicious'])}**")
  lines.append(f"- **缺失举证材料**：**{sum(1 for r in rows if r['missing'])}** 个图斑"
               + ("（离线无法判定，需读系统）" if not any(r['missing'] for r in rows) else ""))
  lines.append(f"- 举证待核（离线）：**{sum(1 for r in rows if r['pending'])}** 个图斑")
  lines.append(f"- 面积有差异：**{sum(1 for r in rows if r['areaDiffs'])}**")
  lines.append("")
  lines.append("| 审核建议 | 数量 |")
  lines.append("|---|---|")
  import collections
  for k, n in collections.Counter(r["suggest"] for r in rows).most_common():
    lines.append(f"| {k} | {n} |")
  lines.append("")
  lines.append(f"面积核对：可比对 {s['comparableCount']} / 差异 {s['fieldDiffCount']} 处"
               f" / 公式校验 可执行 {s['formulaCheckedCount']} 跳过 {s['formulaSkippedCount']}")
  return "\n".join(lines)


def main():
  ap = argparse.ArgumentParser(description="审核意见报告（判定+举证+面积）→ Excel")
  ap.add_argument("--filled", required=True, help="图斑列表导出表")
  ap.add_argument("--truth", required=True, help="套合结果表")
  ap.add_argument("--rules", default=os.path.join(HERE, "..", "rules", "rule-dictionary.json"))
  ap.add_argument("--evidence", help="系统举证状态 JSON（在线读时提供；离线不传）")
  ap.add_argument("--tolerance", type=float, default=0.01)
  ap.add_argument("--out", required=True, help="输出 Excel 路径")
  args = ap.parse_args()

  evidence = None
  if args.evidence and os.path.exists(args.evidence):
    with open(args.evidence, encoding="utf-8") as f:
      evidence = json.load(f)

  result = build_audit(args.filled, args.truth, args.rules, evidence, args.tolerance)
  print(render_summary(result))
  write_excel(result, args.out)
  print(f"\n已写出 → {args.out}")


if __name__ == "__main__":
  main()
