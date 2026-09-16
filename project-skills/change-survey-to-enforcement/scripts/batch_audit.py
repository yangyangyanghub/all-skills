# 批量审核：系统待审核清单 + 套合结果 Excel → 审核意见表
# 用法：python batch_audit.py <list-all.json> <out.xlsx>
import collections
import json
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from lookup_area import lookup, DEFAULT_TRUTH  # noqa: E402

# 地类 → 规则判定（列表级粗判，与 judge.py 决策树一致）
DILEI_RULE = {
  "0702": ("其他", "农村宅基地"),
  "1006": ("农村道路", ""),
  "1202": ("设施农用地", ""),
  "0601": ("非农违法", "（工业用地，需核合法手续）"),
  "0602C": ("非农违法", "（采矿用地，需核采矿权）"),
  "0102": ("非农违法", "（耕地）"),
  "05H1": ("非农违法", "（商业服务业设施用地）"),
  "0701": ("非农违法", "（城镇住宅用地）"),
  "09": ("其他", "（特殊用地）"),
  "0508": ("其他", ""),
}

HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10)
DANGER = PatternFill("solid", fgColor="FFC7CE")
WARN = PatternFill("solid", fgColor="FFEB9C")
OK = PatternFill("solid", fgColor="C6EFCE")


def _style(ws, ncol):
  for c in range(1, ncol + 1):
    cell = ws.cell(row=1, column=c)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
  ws.freeze_panes = "A2"


def main():
  list_file, out_file = sys.argv[1], sys.argv[2]
  rows = json.load(open(list_file, encoding="utf-8"))["rows"]
  res = lookup(DEFAULT_TRUTH, [r["id"] for r in rows])
  by_id = {x["id"]: x for x in res}

  results = []
  for r in rows:
    dl = (r.get("dileiCode") or "").strip()
    op = (r.get("opinion") or "").strip()
    rule = DILEI_RULE.get(dl, ("存疑", "（无对应规则）"))
    consistent = (op == rule[0]) if op else None
    e = by_id.get(r["id"], {})
    in_excel = e.get("found", False)
    ta = e.get("truth", {}).get("areaMu") if in_excel else None
    try:
      fa = float(r.get("areaMu")) if r.get("areaMu") not in (None, "") else None
    except (TypeError, ValueError):
      fa = None
    diff = round(fa - ta, 4) if (fa is not None and ta is not None) else None

    reviews = []
    if op and consistent is False:
      reviews.append(f"判定矛盾：地类{dl or '空'} 但填「{op}」")
    if dl == "":
      reviews.append("实地地类为空")
    if op in ("合法", "设施农用地") and not in_excel:
      reviews.append(f"{op}需核举证（且不在套合表）")
    if op in ("合法",) :
      reviews.append("合法类需核批文覆盖率")
    if diff is not None and abs(diff) > 0.01:
      reviews.append(f"面积差异 {diff} 亩")
    if not in_excel:
      reviews.append("不在套合表（比快照新，需重新导出）")

    if not reviews:
      suggest = "通过" if (consistent is not False) else "需复核"
      level = "OK" if (consistent is not False) else "WARN"
    elif consistent is False and op in ("合法", "设施农用地"):
      suggest = "人工复核"; level = "DANGER"
    else:
      suggest = "需复核"; level = "WARN"

    results.append({
      "id": r["id"], "county": r.get("county") or "", "dilei": dl,
      "opinion": op, "rule": rule[0] + rule[1], "consistent": consistent,
      "inExcel": in_excel, "areaFill": fa, "areaTruth": ta, "areaDiff": diff,
      "reviews": reviews, "suggest": suggest, "level": level, "status": r.get("status") or "",
    })

  # 写 Excel
  wb = openpyxl.Workbook()
  ws = wb.active
  ws.title = "审核意见"
  ws.append(["序号", "线索编号", "县区", "实地地类", "县级填报", "规则判定", "一致性",
             "填报面积", "套合面积", "面积差", "状态", "审核建议", "需人工核的原因"])
  for i, x in enumerate(results, 1):
    ws.append([i, x["id"], x["county"], x["dilei"], x["opinion"], x["rule"],
               {True: "一致", False: "不一致", None: "未填报"}[x["consistent"]],
               x["areaFill"] if x["areaFill"] is not None else "",
               x["areaTruth"] if x["areaTruth"] is not None else ("无" if x["inExcel"] else ""),
               x["areaDiff"] if x["areaDiff"] is not None else "",
               x["status"], x["suggest"], "；".join(x["reviews"])])
    row = ws.max_row
    if x["level"] == "DANGER": ws.cell(row=row, column=7).fill = DANGER
    elif x["level"] == "WARN": ws.cell(row=row, column=7).fill = WARN
    elif x["level"] == "OK": ws.cell(row=row, column=7).fill = OK
    if x["reviews"]: ws.cell(row=row, column=13).fill = WARN if x["level"] != "DANGER" else DANGER
  _style(ws, 13)
  for i, w in enumerate([5, 24, 10, 10, 12, 24, 9, 9, 9, 9, 12, 12, 40], 1):
    ws.column_dimensions[get_column_letter(i)].width = w

  # 汇总
  review = [x for x in results if x["reviews"]]
  print("=" * 70)
  print(f"待审核图斑：{len(results)} 个")
  print(f"  与填报不一致：{sum(1 for x in results if x['consistent'] is False)}")
  print(f"  需人工复核：{len(review)} 个")
  print(f"  面积差异：{sum(1 for x in results if x['areaDiff'] is not None and abs(x['areaDiff']) > 0.01)}")
  print(f"  在套合表：{sum(1 for x in results if x['inExcel'])} / 不在：{sum(1 for x in results if not x['inExcel'])}")
  print("=" * 70)
  print("\n【需人工复核清单】")
  for x in review:
    print(f"  {x['id']:<24} 地类{x['dilei'] or '(空)':<6} 填报{x['opinion']:<10} → {x['suggest']}")
    for w_ in x["reviews"]:
      print(f"      · {w_}")

  wb.save(out_file)
  print(f"\n已写出 → {out_file}")


if __name__ == "__main__":
  main()
